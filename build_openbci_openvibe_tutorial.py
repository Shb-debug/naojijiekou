from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(r"E:\dijikeji\naojijiekou")
IMG_DIR = Path(r"C:\Users\lenovo\Desktop\图片")
OUT = ROOT / "OpenBCI_OpenViBE_环境配置与使用教程_聊天记录版.docx"
FONT = "Microsoft YaHei"
MONO = "Consolas"
NAVY = "17365D"
BLUE = "2E74B5"
MUTED = "687386"
GREEN = "C6EFCE"
GREEN_DARK = "216E39"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F9"
LINE = "D8E0EA"
WHITE = "FFFFFF"


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in kwargs[edge]:
                    element.set(qn("w:" + key), str(kwargs[edge][key]))


def set_paragraph_shading(p, fill):
    p_pr = p._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(p, edge="bottom", color=LINE, size="8", space="2"):
    p_pr = p._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    el = borders.find(qn(f"w:{edge}"))
    if el is None:
        el = OxmlElement(f"w:{edge}")
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), space)
    el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_keep(p, with_next=False, together=False):
    p.paragraph_format.keep_with_next = with_next
    p.paragraph_format.keep_together = together


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=MUTED)


def style_paragraph(p, size=10.5, color="222222", bold=False, align=None, before=0, after=6, line=1.25):
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for run in p.runs:
        set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_body(doc, text, bold_prefix=None, after=6):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, color="222222", bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5, color="222222")
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color="222222")
    style_paragraph(p, after=after)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="222222")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="222222")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, size=16, color=BLUE, bold=True)
    elif level == 2:
        set_run_font(r, size=13, color=BLUE, bold=True)
    else:
        set_run_font(r, size=12, color=NAVY, bold=True)
    set_keep(p, with_next=True)
    return p


def add_chat(doc, text, img_num=None, caption=None, width=None, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.left_indent = Inches(1.42)
    p.paragraph_format.right_indent = Inches(0.02)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6 if img_num else 10)
    p.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(p, GREEN)
    set_paragraph_border(p, "top", color="B4DDBB", size="4", space="1")
    set_paragraph_border(p, "bottom", color="B4DDBB", size="4", space="1")
    set_paragraph_border(p, "left", color="B4DDBB", size="4", space="1")
    set_paragraph_border(p, "right", color="B4DDBB", size="4", space="1")
    run = p.add_run("我  ·  ")
    set_run_font(run, size=9, color=GREEN_DARK, bold=True)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color="1F2933")
    set_keep(p, with_next=bool(img_num))
    if img_num:
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        img_p.paragraph_format.space_before = Pt(0)
        img_p.paragraph_format.space_after = Pt(2)
        img_p.paragraph_format.keep_together = True
        img_path = IMG_DIR / f"{img_num}.png"
        run = img_p.add_run()
        run.add_picture(str(img_path), width=Inches(width or 6.15))
        drawing = run._r.xpath(".//wp:docPr")
        if drawing:
            drawing[0].set("descr", f"OpenBCI/OpenViBE 操作示意图 {img_num}")
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cap.paragraph_format.space_before = Pt(0)
        cap.paragraph_format.space_after = Pt(4)
        cr = cap.add_run(caption or f"图 {img_num} · 操作示意")
        set_run_font(cr, size=8.5, color=MUTED, italic=True)
    if note:
        np = doc.add_paragraph()
        np.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        np.paragraph_format.space_after = Pt(6)
        nr = np.add_run(note)
        set_run_font(nr, size=9, color=MUTED, italic=True)


def add_section_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text.upper())
    set_run_font(r, size=9, color=GREEN_DARK, bold=True)
    set_paragraph_border(p, "bottom", color=GREEN_DARK, size="10", space="3")
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    set_paragraph_shading(p, "EEF2F6")
    set_paragraph_border(p, "left", color=BLUE, size="18", space="4")
    r = p.add_run(text)
    set_run_font(r, name=MONO, size=9.5, color=NAVY)
    return p


def add_info_box(doc, title, body, fill=LIGHT_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(p, fill)
    set_paragraph_border(p, "left", color=BLUE, size="18", space="5")
    r = p.add_run(title + "\n")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(body)
    set_run_font(r, size=10, color="334155")
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ]:
        st = styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.15

    # Running furniture
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("配置教程  /  OpenBCI × OpenViBE")
    set_run_font(hr, size=8.5, color=MUTED, bold=True)
    fp = sec.footer.paragraphs[0]
    add_page_number(fp)

    # Title block
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("图文操作手册 · CHAT LOG EDITION")
    set_run_font(r, size=9, color=GREEN_DARK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("OpenBCI + OpenViBE\n环境配置与使用教程")
    set_run_font(r, size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run("从 OpenBCI GUI 采集脑电，经 LSL 传输到 OpenViBE Acquisition Server，再进入 Designer 场景。")
    set_run_font(r, size=12, color=MUTED)

    meta = doc.add_table(rows=4, cols=4)
    set_table_geometry(meta, [1350, 3330, 1350, 3330], indent_dxa=120)
    header = meta.rows[0]
    hdr_pr = header._tr.get_or_add_trPr()
    hdr_flag = OxmlElement("w:tblHeader")
    hdr_flag.set(qn("w:val"), "true")
    hdr_pr.append(hdr_flag)
    for i, val in enumerate(["项目", "配置", "项目", "配置"]):
        cell = header.cells[i]
        set_cell_shading(cell, "DCE6F1")
        set_cell_border(cell, top={"val":"single","sz":"4","color":LINE}, bottom={"val":"single","sz":"4","color":LINE}, left={"val":"single","sz":"4","color":LINE}, right={"val":"single","sz":"4","color":LINE})
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        rr = cell.paragraphs[0].add_run(val)
        set_run_font(rr, size=9, color=NAVY, bold=True)
    meta_entries = [
        ("适用系统", "Windows 10/11", "配图顺序", "1 → 12"),
        ("数据链路", "OpenBCI GUI → LSL → Acquisition Server → Designer", "建议版本", "以官网下载页为准"),
        ("输出目标", "实时显示、记录，并供 OpenViBE 场景继续处理", "文档性质", "截图参考 + 可执行步骤"),
    ]
    for row, vals in zip(meta.rows[1:], meta_entries):
        for i, val in enumerate(vals):
            cell = row.cells[i]
            set_cell_shading(cell, LIGHT_GRAY if i % 2 == 0 else WHITE)
            set_cell_border(cell, top={"val":"single","sz":"4","color":LINE}, bottom={"val":"single","sz":"4","color":LINE}, left={"val":"single","sz":"4","color":LINE}, right={"val":"single","sz":"4","color":LINE})
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            rr = cell.paragraphs[0].add_run(val)
            set_run_font(rr, size=9 if i % 2 == 0 else 9.5, color=NAVY if i % 2 == 0 else "334155", bold=(i % 2 == 0))

    add_info_box(doc, "先记住这一条主线", "OpenBCI GUI 负责连接硬件并产生 LSL 流；OpenViBE Acquisition Server 负责接收 LSL；OpenViBE Designer 负责打开并运行 XML 场景。三个程序都打开时，启动顺序建议保持：GUI → LSL → Acquisition Server → Designer。", fill="EAF6EE")
    add_body(doc, "说明：聊天气泡中的文字是本教程整理后的操作指令；截图是用户提供的参考示意。截图里显示的版本/界面可能与当前下载版本略有差异，若按钮位置变化，以软件当前界面中含义相同的控件为准。")

    add_heading(doc, "1. 下载与安装", 1)
    add_body(doc, "建议先把两个软件安装包下载到同一台用于采集的电脑上。OpenBCI GUI 可作为独立应用运行；OpenViBE 需要使用 Acquisition Server 接收数据，并在 Designer 中加载场景。")
    p = doc.add_paragraph(style="List Number")
    r = p.add_run("下载 OpenViBE：")
    set_run_font(r, size=10.5, color="222222", bold=True)
    add_hyperlink(p, "OpenViBE 官方网站", "https://openvibe.inria.fr/")
    p.add_run("（进入官网后选择适合当前系统的下载项。）")
    style_paragraph(p, after=4)
    p = doc.add_paragraph(style="List Number")
    r = p.add_run("下载 OpenBCI GUI：")
    set_run_font(r, size=10.5, color="222222", bold=True)
    add_hyperlink(p, "OpenBCI GUI 官方文档/下载入口", "https://docs.openbci.com/Software/OpenBCISoftware/GUIDocs/")
    p.add_run("。Windows 用户下载后解压，进入 OpenBCI_GUI 文件夹，双击 OpenBCI_GUI.exe。")
    style_paragraph(p, after=4)
    add_body(doc, "如果使用 Cyton 或 Cyton+Daisy，Windows 可能需要先安装最新 FTDI 驱动；如果使用 Ganglion，则按硬件说明准备 USB Dongle/BLE 连接。")
    add_body(doc, "建议把 OpenViBE 的版本、OpenBCI GUI 的版本和项目 XML 场景记录下来，后续遇到连接问题时更容易定位。")

    add_heading(doc, "2. OpenBCI GUI：连接硬件并确认脑电波形", 1)
    add_section_label(doc, "聊天记录 · GUI 启动与采集")
    add_chat(doc, "下载后打开 eeg 开头的程序；电脑先连接网络。", 1, "图 1 · 先启动与设备配套的程序，并确认电脑网络正常。", width=4.9)
    add_chat(doc, "打开 OpenBCI GUI。第一次使用时，先确认硬件已经通电、USB Dongle/串口连接正常。", 2, "图 2 · OpenBCI GUI 的 System Control Panel。", width=6.15)
    add_chat(doc, "如图设置：Data Source 选择与你的硬件对应的 live 选项；传输协议按实际连接方式选择 Serial（from Dongle）或 WiFi（from WiFi Shield）；通道数和采样率必须与硬件/实验设置一致。", None, note="截图示例为 16 channels、500 Hz；如果你的板卡或实验设置不同，不要机械照抄这两个值。")
    add_chat(doc, "设置好后，点击 Start Session。", 3, "图 3 · Start Session 后进入数据界面。", width=6.15)
    add_chat(doc, "进入后点击右上角的 Start Data Stream。", None, note="先观察界面是否从静止状态进入实时刷新；如果没有刷新，先回到 Data Source、串口/无线连接和电极接触处排查。")
    add_chat(doc, "就可以看到脑电波了。确认左侧通道曲线会变化，再继续配置 LSL。", 4, "图 4 · OpenBCI GUI 已开始实时显示 EEG 波形。", width=6.15)

    add_info_box(doc, "信号检查小口诀", "先看有没有曲线，再看是否持续刷新，最后再看是否存在明显饱和、整屏噪声或大量断线。不要在尚未确认原始波形正常时直接排查 OpenViBE。", fill="FFF8E8")

    add_heading(doc, "3. OpenBCI GUI：通过 Networking 输出 LSL", 1)
    add_section_label(doc, "聊天记录 · LSL 输出")
    add_chat(doc, "在任意一个数据显示窗口的左上角下拉框中选择 Networking。", 5, "图 5 · 从窗口类型下拉框进入 Networking。", width=5.5)
    add_chat(doc, "打开后，先看右上角 Protocol。截图中的 UDP 是可选协议示例；本教程连接 OpenViBE 时要把它切换为 LSL。", 6, "图 6 · Protocol 下拉框中可以看到 UDP、LSL、OSC、Serial 等选项。", width=5.55)
    add_chat(doc, "这个窗口的右上角选 LSL。", 7, "图 7 · 切换到 LSL 协议后的界面。", width=5.55)
    add_chat(doc, "这里选 TimeSeriesRaw；Stream 1 的 Name 填 obci_eeg1，Type 填 EEG，然后点击 Start LSL Stream。", 8, "图 8 · LSL 输出示例：Stream 1 使用 TimeSeriesRaw / obci_eeg1 / EEG。", width=5.55)
    add_body(doc, "推荐只启用一个 EEG 流，避免后续在 Acquisition Server 中选错流。其他 Stream 可保持 None；如果你确实要同时输出多个流，请给每个流设置不同且易识别的 Name。")
    add_code(doc, "LSL Stream name: obci_eeg1\nLSL Stream type: EEG\n本机优先：127.0.0.1（同一台电脑时）")

    add_heading(doc, "4. OpenViBE Acquisition Server：接收 LSL", 1)
    add_section_label(doc, "聊天记录 · Acquisition Server")
    add_chat(doc, "设置完成后，直接点击 Start（在 OpenBCI GUI 中保持 LSL 流运行）。然后打开 OpenViBE Acquisition Server。", 9, "图 9 · OpenViBE Acquisition Server 初始界面。", width=6.15)
    add_chat(doc, "如图改好 Driver、Port、block：Driver 选择 LabStreamingLayer (LSL)；Connection port 保持 1024；Sample count per sent block 建议先用 32。", 10, "图 10 · 选择 LSL 驱动并准备打开 Driver Properties。", width=6.15)
    add_chat(doc, "点击 Driver Properties。", 11, "图 11 · LSL Device configuration。重点是让 Signal stream 指向 obci_eeg1 / openbcigui；Marker stream 没有标记流时保持 None。", width=3.7)
    add_body(doc, "在 Device configuration 中建议这样核对：Identifier 可保持默认；Age/Gender 只是元数据；Fallback Sampling Frequency 留空时优先使用 LSL 流自带采样率，如果设备没有提供有效采样率，再填入与你的 OpenBCI GUI 一致的数值。设置完成点击 Apply。")
    add_chat(doc, "之后点击依次 Connect、Play。看到设备状态开始更新，说明 Acquisition Server 已经开始向 OpenViBE Designer 提供数据。", None, note="若 Connect 失败，先确认 OpenBCI GUI 中的 Start LSL Stream 仍处于运行状态，并检查 obci_eeg1 名称是否完全一致。")

    add_heading(doc, "5. OpenViBE Designer：打开 XML 场景并运行", 1)
    add_section_label(doc, "聊天记录 · Designer")
    add_chat(doc, "打开 OpenViBE Designer。", 12, "图 12 · OpenViBE Designer：打开 XML 场景并点击运行按钮。", width=6.15)
    add_chat(doc, "打开 .xml 文件后运行即可。点击 File → Open，选择你的场景文件，再点击工具栏上的三角形 Run 按钮。", None, note="当前项目中可见的示例场景包括 CARGAME_OpenViBE_SSVEP.xml 和 SUPER-mario.xml；请按实际实验目标选择，并核对其中的通道选择器、滤波范围、分段参数和脚本路径。")
    add_info_box(doc, "项目场景核对", "如果使用 CARGAME_OpenViBE_SSVEP.xml，场景默认会选择后部通道 2;4;6，并使用 4–32 Hz、1.5 秒分段；如果使用 SUPER-mario.xml，场景默认选择 1;2;15;14;16;13，并使用 1–30 Hz、0.32 秒分段。通道编号必须按你的实际电极排列核对，不能只按文件名猜测。", fill="F0F4FA")

    add_heading(doc, "6. 一次完整运行顺序", 1)
    add_body(doc, "每次实验都可以按下面顺序启动，便于定位是哪一环没有工作：")
    for t in [
        "打开 OpenBCI GUI，选择正确硬件与数据源。",
        "Start Session → Start Data Stream，确认 GUI 中出现实时波形。",
        "在 Networking 中选择 LSL，Stream 1 设置为 obci_eeg1 / EEG，点击 Start LSL Stream。",
        "打开 Acquisition Server，Driver 选 LabStreamingLayer (LSL)，端口 1024，block 32；点击 Driver Properties → Apply → Connect → Play。",
        "打开 OpenViBE Designer，载入 XML 场景，点击 Run。",
        "观察 Designer 中的 Signal Display / Spectrum Display 或下游处理模块是否有输入。",
    ]:
        add_number(doc, t)

    add_heading(doc, "7. 常见问题快速排查", 1)
    problems = [
        ("GUI 没有波形", "检查板卡是否通电、串口/Dongle 是否识别、Data Source 是否选对、电极接触和阻抗是否正常；先在 GUI 内解决，再继续 LSL。"),
        ("Networking 里找不到 LSL", "确认窗口类型已经切到 Networking，Protocol 选择 LSL，Stream 1 不是 None，Name 为 obci_eeg1，Type 为 EEG，并点击 Start LSL Stream。"),
        ("Acquisition Server 连接失败", "确认 Driver 为 LabStreamingLayer (LSL)，端口使用 1024；同一台电脑优先使用 127.0.0.1；检查防火墙和流名称是否一致。"),
        ("Designer 没有数据", "确认 Acquisition Server 已 Connect + Play；XML 场景中的 Acquisition client 主机/端口与 Server 一致；Channel Selector 的编号要与硬件电极排列一致。"),
        ("波形明显饱和或噪声很大", "降低电极阻抗、检查参考/地电极、远离电源和 USB 干扰；不要把饱和波形误认为有效脑电。"),
    ]
    for title, body in problems:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(title + "：")
        set_run_font(r, size=10.5, color=NAVY, bold=True)
        r = p.add_run(body)
        set_run_font(r, size=10.5, color="222222")

    add_heading(doc, "8. 记录与安全提示", 1)
    add_body(doc, "OpenBCI GUI 默认会把会话和原始记录保存到用户 Documents/OpenBCI_GUI 目录下；具体子目录和文件格式会随 GUI 版本、保存设置而变化。实验完成后建议同时保存：GUI 版本、OpenViBE 版本、XML 场景、LSL stream 名称、采样率、通道映射和实验日期。")
    add_info_box(doc, "安全提示", "本教程只描述软件连接与数据流，不构成医疗建议。若实验包含闪烁视觉刺激，请预先评估光敏性癫痫等风险，准备人工停止方式，并在正式实验前用低强度、短时长进行测试。", fill="FDECEC")

    add_heading(doc, "附：官方入口", 1)
    p = doc.add_paragraph()
    r = p.add_run("OpenViBE：")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    add_hyperlink(p, "OpenViBE 官网", "https://openvibe.inria.fr/")
    style_paragraph(p, after=4)
    p = doc.add_paragraph()
    r = p.add_run("OpenBCI GUI 文档：")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    add_hyperlink(p, "OpenBCI GUI 官方文档", "https://docs.openbci.com/Software/OpenBCISoftware/GUIDocs/")
    style_paragraph(p, after=4)
    add_body(doc, "官方页面会随软件版本更新；下载前以页面显示的当前版本和系统要求为准。")

    # Core properties are deliberately generic to avoid leaking local account names.
    doc.core_properties.title = "OpenBCI + OpenViBE 环境配置与使用教程"
    doc.core_properties.subject = "LSL 数据链路配置与聊天记录式图文操作"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
